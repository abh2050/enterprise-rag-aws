"""Deterministic SYNTHETIC binary fixtures (PDF/DOCX) generated at test/eval time.

Nothing here is real data. Generated files are written to temporary directories and never committed.
"""

from __future__ import annotations

import io
from pathlib import Path

EICAR_TEXT = "X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"


def pdf_with_table() -> bytes:
    """Two-page text PDF: headings, paragraphs and a ruled table on page 2."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, invariant=1)
    styles = getSampleStyleSheet()
    table = Table(
        [
            ["Region", "Shipping fee (USD)", "Delivery days"],
            ["North", "12", "3"],
            ["South", "15", "4"],
            ["West", "18", "5"],
        ],
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build(
        [
            Paragraph("Logistics Handbook", styles["Title"]),
            Paragraph("Shipping Overview", styles["Heading1"]),
            Paragraph(
                "All shipments are dispatched from the central warehouse within one business day of the "
                "order being confirmed. Fragile items are packed in double-walled cartons.",
                styles["BodyText"],
            ),
            PageBreak(),
            Paragraph("Regional Shipping Fees", styles["Heading1"]),
            Paragraph(
                "The table below lists shipping fees and delivery times by region.", styles["BodyText"]
            ),
            table,
        ]
    )
    return buf.getvalue()


def scanned_pdf() -> bytes:
    """One text page + one image-only page (simulates a scanned page with no text layer)."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (800, 300), "white")
    ImageDraw.Draw(img).text((20, 120), "Scanned page: warranty period is 24 months.", fill="black")
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter, invariant=1)
    c.setFont("Helvetica", 11)
    c.drawString(72, 700, "Warranty Terms. This first page has a normal text layer describing coverage.")
    c.showPage()
    c.drawImage(ImageReader(img_buf), 72, 400, width=450, height=170)
    c.showPage()
    c.save()
    return buf.getvalue()


def encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(pdf_with_table()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password="synthetic-user-pw", owner_password="synthetic-owner-pw")  # noqa: S106
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_with_table() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Security Awareness Policy", level=1)
    d.add_paragraph("All staff must complete security awareness training within 30 days of joining.")
    d.add_heading("Password Rules", level=2)
    d.add_paragraph("Passwords must be at least 14 characters and must not be reused across systems.")
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate(
        [("Control", "Frequency"), ("Phishing simulation", "Quarterly"), ("Access review", "Semi-annual")]
    ):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def write_collection(root: Path, name: str, files: dict[str, bytes | str], access_yaml: str) -> Path:
    folder = root / name
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "_access.yaml").write_text(access_yaml)
    for filename, data in files.items():
        path = folder / filename
        path.write_bytes(data.encode() if isinstance(data, str) else data)
    return folder
